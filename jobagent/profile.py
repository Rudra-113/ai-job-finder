"""Your CV and saved facts. Answers may use only what is stored here."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from jobagent.paths import atomic_write, file_lock, read_json

_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_URL = re.compile(r"https?://", re.I)


@dataclass
class Experience:
    title: str
    company: str
    start: str
    end: str
    details: str


@dataclass
class Profile:
    full_name: str = ""
    email: str = ""
    phone: str = ""
    city: str = ""
    country: str = ""
    portfolio: str = ""
    linkedin: str = ""
    github: str = ""
    work_authorized: str = ""
    needs_sponsorship: str = ""
    skills: list[str] = field(default_factory=list)
    education_raw: str = ""
    experience_raw: str = ""
    projects: str = ""
    extra_facts: str = ""
    cv_path: str = ""
    cv_text: str = ""

    @property
    def first_name(self) -> str:
        parts = self.full_name.split()
        return parts[0] if parts else ""

    @property
    def last_name(self) -> str:
        parts = self.full_name.split()
        return " ".join(parts[1:]) if len(parts) > 1 else ""

    def problem(self) -> str | None:
        if not self.full_name.strip() or " " not in self.full_name.strip():
            return "Enter your full name, including a first and last name, on the Profile page."
        if not _EMAIL.match(self.email.strip()):
            return "Enter an email address on the Profile page."
        if not self.cv_path or not Path(self.cv_path).is_file():
            return "Upload your CV on the Profile page."
        if not (self.cv_text.strip() or self.skills or self.experience_raw.strip()):
            return (
                "The CV had no readable text. Upload a text-based PDF, DOCX, or TXT file, "
                "or add skills and experience on the Profile page."
            )
        return None

    def to_dict(self) -> dict:
        return {
            "full_name": self.full_name,
            "email": self.email,
            "phone": self.phone,
            "city": self.city,
            "country": self.country,
            "portfolio": self.portfolio,
            "linkedin": self.linkedin,
            "github": self.github,
            "work_authorized": self.work_authorized,
            "needs_sponsorship": self.needs_sponsorship,
            "skills": self.skills,
            "education_raw": self.education_raw,
            "experience_raw": self.experience_raw,
            "projects": self.projects,
            "extra_facts": self.extra_facts,
            "cv_path": self.cv_path,
            "cv_text": self.cv_text,
        }

    @classmethod
    def from_dict(cls, data: dict) -> "Profile":
        skills = data.get("skills") or []
        if isinstance(skills, str):
            skills = [part.strip() for part in skills.split(",") if part.strip()]
        return cls(
            full_name=str(data.get("full_name") or ""),
            email=str(data.get("email") or ""),
            phone=str(data.get("phone") or ""),
            city=str(data.get("city") or ""),
            country=str(data.get("country") or ""),
            portfolio=str(data.get("portfolio") or ""),
            linkedin=str(data.get("linkedin") or ""),
            github=str(data.get("github") or ""),
            work_authorized=str(data.get("work_authorized") or ""),
            needs_sponsorship=str(data.get("needs_sponsorship") or ""),
            skills=[str(item).strip() for item in skills if str(item).strip()],
            education_raw=str(data.get("education_raw") or ""),
            experience_raw=str(data.get("experience_raw") or ""),
            projects=str(data.get("projects") or ""),
            extra_facts=str(data.get("extra_facts") or ""),
            cv_path=str(data.get("cv_path") or ""),
            cv_text=str(data.get("cv_text") or ""),
        )


def profile_path(home: Path) -> Path:
    return home / "profile.json"


def load_profile(home: Path) -> Profile:
    path = profile_path(home)
    with file_lock(path):
        return Profile.from_dict(read_json(path, {}))


def save_profile(home: Path, profile: Profile) -> None:
    path = profile_path(home)
    with file_lock(path):
        atomic_write(path, profile.to_dict())


def parse_experience(raw: str) -> list[Experience]:
    rows, _bad = parse_pipe_rows(raw, 4)
    return [
        Experience(parts[0], parts[1], parts[2], parts[3], parts[4] if len(parts) > 4 else "")
        for parts in rows
    ]


def parse_pipe_rows(raw: str, expected: int) -> tuple[list[list[str]], list[str]]:
    rows: list[list[str]] = []
    bad: list[str] = []
    for lineno, line in enumerate((raw or "").splitlines(), start=1):
        stripped = line.strip()
        if not stripped or stripped.startswith("#"):
            continue
        parts = [part.strip() for part in stripped.split("|")]
        if len(parts) < expected or not any(parts):
            bad.append(f"Line {lineno} needs at least {expected} fields separated by |")
            continue
        rows.append(parts)
    return rows, bad


def extract_cv_text(path: Path) -> str:
    suffix = path.suffix.casefold()
    try:
        if suffix in {".txt", ".md"}:
            return path.read_text(encoding="utf-8", errors="replace")[:100_000]
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            parts = [(page.extract_text() or "") for page in reader.pages]
            return "\n".join(parts)[:100_000]
        if suffix == ".docx":
            from docx import Document

            document = Document(str(path))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" ".join(cell.text for cell in row.cells))
            return "\n".join(parts)[:100_000]
    except Exception:
        return ""
    return ""


def clean_optional_url(value: str, label: str) -> tuple[str, str | None]:
    text = (value or "").strip()
    if not text:
        return "", None
    if not _URL.match(text):
        return "", f"{label} must start with http:// or https://"
    return text, None


def digits(value: str) -> str:
    return re.sub(r"\D", "", value or "")


def validate_profile_fields(form) -> tuple[dict, str | None, list[str]]:
    name = re.sub(r"\s+", " ", (form.get("full_name") or "").strip())
    email = (form.get("email") or "").strip()
    phone = (form.get("phone") or "").strip()
    if name and " " not in name:
        return {}, "Enter a first and last name.", []
    if email and not _EMAIL.match(email):
        return {}, "Enter a valid email address.", []
    if phone and len(digits(phone)) < 7:
        return {}, "Enter a phone number with at least 7 digits, or leave it blank.", []
    portfolio, error = clean_optional_url(form.get("portfolio") or "", "Portfolio URL")
    if error:
        return {}, error, []
    linkedin, error = clean_optional_url(form.get("linkedin") or "", "LinkedIn URL")
    if error:
        return {}, error, []
    github, error = clean_optional_url(form.get("github") or "", "GitHub URL")
    if error:
        return {}, error, []
    work = (form.get("work_authorized") or "").strip()
    sponsor = (form.get("needs_sponsorship") or "").strip()
    if work not in {"", "yes", "no"}:
        return {}, "Choose a work-authorization answer.", []
    if sponsor not in {"", "yes", "no"}:
        return {}, "Choose a sponsorship answer.", []
    education = form.get("education_raw") or ""
    experience = form.get("experience_raw") or ""
    _edu_rows, edu_bad = parse_pipe_rows(education, 4)
    _exp_rows, exp_bad = parse_pipe_rows(experience, 4)
    skills = [part.strip() for part in (form.get("skills") or "").split(",") if part.strip()]
    payload = {
        "full_name": name,
        "email": email,
        "phone": phone,
        "city": (form.get("city") or "").strip(),
        "country": (form.get("country") or "").strip(),
        "portfolio": portfolio,
        "linkedin": linkedin,
        "github": github,
        "work_authorized": work,
        "needs_sponsorship": sponsor,
        "skills": skills,
        "education_raw": education.strip(),
        "experience_raw": experience.strip(),
        "projects": (form.get("projects") or "").strip(),
        "extra_facts": (form.get("extra_facts") or "").strip(),
    }
    return payload, None, edu_bad + exp_bad
